from pathlib import Path
import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs/监控与语义识别及提醒机制说明.md'
OUTPUT = SOURCE.with_suffix('.docx')
FIGURE = ROOT / 'docs/assets/monitoring-mechanism-flow.png'
NAVY = '17365A'
BLUE = '155E75'


def font(size):
    return ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', size)


def diagram():
    FIGURE.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (1800, 790), '#ffffff')
    draw = ImageDraw.Draw(image)
    boxes = [(50, 45, '精确目标群窗口', '微信在线 · 独立窗口就绪'),
             (660, 45, '原生采集与转发补全', 'UIA事件唤醒 + 定期核对'),
             (1270, 45, '本地规则识别', '摘要 · 报名截止 · 活动时间'),
             (1270, 305, '活动 / 待核实线索', '日程能否确定，分别处理'),
             (660, 305, 'SQLite 去重与版本', '关联原活动 · 旧任务作废'),
             (50, 305, 'outbox 提醒队列', '三项通知 · 活动开始前提醒'),
             (50, 565, '到期任务检查', '目标间隔5秒 · 考虑故障退避'),
             (660, 565, '检查发送条件', '发送模式与本机邮箱凭据'),
             (1270, 565, 'QQ SMTP / 预览保留', '就绪发送；未就绪继续排队')]
    for x, y, title, subtitle in boxes:
        draw.rounded_rectangle((x, y, x + 480, y + 155), radius=22, fill='#F0F6FA', outline='#AAC6D7', width=3)
        draw.text((x+240, y+52), title, font=font(32), anchor='mm', fill='#17365A')
        draw.text((x+240, y+109), subtitle, font=font(25), anchor='mm', fill='#426073')
    for start, end in [((530,122),(660,122)),((1140,122),(1270,122)),((1510,200),(1510,305)),
                       ((1270,382),(1140,382)),((660,382),(530,382)),((290,460),(290,565)),
                       ((530,642),(660,642)),((1140,642),(1270,642))]:
        draw.line((start,end),fill='#31748B',width=6)
        x,y=end
        if end[0]>start[0]:triangle=[(x,y),(x-18,y-11),(x-18,y+11)]
        elif end[0]<start[0]:triangle=[(x,y),(x+18,y-11),(x+18,y+11)]
        else:triangle=[(x,y),(x-11,y-18),(x+11,y-18)]
        draw.polygon(triangle,fill='#31748B')
    image.save(FIGURE, dpi=(240,240))


def shading(element, fill):
    cell = OxmlElement('w:shd'); cell.set(qn('w:fill'), fill); element.append(cell)


def configure_run(run, size=None, bold=None, color=None, mono=False):
    run.font.name = 'Consolas' if mono else 'Microsoft YaHei'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    if color is not None:run.font.color.rgb=RGBColor.from_string(color)


def hyperlink(paragraph, label, url):
    if not url.startswith(('http://','https://')):
        url='https://github.com/Fui2311/wechat_tracking/blob/main/docs/'+url
    rel=paragraph.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r'); prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),BLUE);prop.append(color)
    run.append(prop);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)


def inline(paragraph, text, size=None, color=None):
    text=text.replace('–','-').replace('‑','-')
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))',text):
        if not token:continue
        link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if link:
            hyperlink(paragraph,*link.groups());continue
        bold=token.startswith('**') and token.endswith('**')
        mono=token.startswith('`') and token.endswith('`')
        value=token[2:-2] if bold else token[1:-1] if mono else token
        run=paragraph.add_run(value);configure_run(run,size,bold,color,mono)


def field(paragraph, code):
    run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    run=paragraph.add_run();instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' '+code+' ';run._r.append(instruction)
    run=paragraph.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)


def build():
    diagram()
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.65);section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
    section.header_distance=Cm(.65);section.footer_distance=Cm(.65)
    normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
    normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(5)
    for name,size,color in [('Title',22,NAVY),('Heading 1',15,NAVY),('Heading 2',12,BLUE)]:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string(color)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'微软雅黑')
        style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(12 if name!='Title' else 0);style.paragraph_format.space_after=Pt(6)
    header=section.header.paragraphs[0];inline(header,'微信群讲座监控  |  方案与机制说明',8,'66788A')
    footer=section.footer.paragraphs[0];footer.alignment=2
    inline(footer,'0.4.0  ·  2026-10-15     第 ',8,'66788A');field(footer,'PAGE');inline(footer,' 页 / 共 ',8,'66788A');field(footer,'NUMPAGES');inline(footer,' 页',8,'66788A')
    props=doc.core_properties;props.title='微信群讲座监控、语义识别与提醒机制说明';props.subject='当前实现、运行快照、规则识别及QQ提醒机制'
    props.author='WeChat Lecture Tracker';props.keywords='微信,讲座,监控,语义识别,QQ邮箱,提醒';props.created=datetime(2026,10,15,tzinfo=timezone.utc)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings._element.append(update)
    lines=SOURCE.read_text(encoding='utf-8').splitlines();index=0
    while index<len(lines):
        line=lines[index]
        if not line.strip():index+=1;continue
        if line.startswith('```'):
            language=line[3:];block=[];index+=1
            while index<len(lines) and not lines[index].startswith('```'):
                block.append(lines[index]);index+=1
            if language=='mermaid':
                paragraph=doc.add_paragraph();paragraph.paragraph_format.keep_together=True
                paragraph.add_run().add_picture(str(FIGURE),width=Cm(17.2))
                caption=doc.add_paragraph();caption.alignment=1
                inline(caption,'图1  当前监控与提醒主链路（发送条件未就绪时保留任务）',9,'66788A')
            else:
                for code_index,value in enumerate(block):
                    paragraph=doc.add_paragraph();paragraph.paragraph_format.left_indent=Cm(.2);paragraph.paragraph_format.space_after=Pt(1);paragraph.paragraph_format.line_spacing=1.05
                    paragraph.paragraph_format.keep_with_next=code_index<len(block)-1
                    shading(paragraph._p.get_or_add_pPr(),'F1F4F7')
                    run=paragraph.add_run(value);configure_run(run,8.8,mono=True)
            index+=1;continue
        if line.startswith('|'):
            rows=[]
            while index<len(lines) and lines[index].startswith('|'):
                parts=[part.strip() for part in lines[index].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',part.replace(' ','')) for part in parts):rows.append(parts)
                index+=1
            count=len(rows[0]);table=doc.add_table(rows=0,cols=count);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            widths=[4.0,13.2] if count==2 else [3.2,3.2,10.8] if count==3 else [5.1,2.8,3.0,6.3]
            for column,width in zip(table.columns,widths):column.width=Cm(width)
            for number,values in enumerate(rows):
                cells=table.add_row().cells
                rowprops=cells[0]._tc.getparent().get_or_add_trPr();split=OxmlElement('w:cantSplit');rowprops.append(split)
                if number==0:
                    repeat=OxmlElement('w:tblHeader');rowprops.append(repeat)
                for column,(cell,value) in enumerate(zip(cells,values)):
                    cell.width=Cm(widths[column]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cellprops=cell._tc.get_or_add_tcPr();shading(cellprops,NAVY if number==0 else 'F1F6FA' if number%2 else 'FFFFFF')
                    paragraph=cell.paragraphs[0];paragraph.paragraph_format.space_before=Pt(3);paragraph.paragraph_format.space_after=Pt(3);paragraph.paragraph_format.line_spacing=1.1
                    paragraph.paragraph_format.keep_with_next=number<len(rows)-1
                    inline(paragraph,value,9.3,'FFFFFF' if number==0 else '23394D')
                    if number==0:
                        for run in paragraph.runs:run.bold=True
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        if line.startswith('# '):paragraph=doc.add_paragraph(style='Title');inline(paragraph,line[2:])
        elif line.startswith('## '):paragraph=doc.add_paragraph(style='Heading 1');inline(paragraph,line[3:])
        elif line.startswith('### '):paragraph=doc.add_paragraph(style='Heading 2');inline(paragraph,line[4:])
        elif line.startswith('- '):paragraph=doc.add_paragraph(style='List Bullet');inline(paragraph,line[2:])
        elif line.startswith('> '):
            paragraph=doc.add_paragraph();paragraph.paragraph_format.left_indent=Cm(.4);paragraph.paragraph_format.space_after=Pt(2)
            shading(paragraph._p.get_or_add_pPr(),'F0F6FA');inline(paragraph,line[2:])
        else:
            paragraph=doc.add_paragraph();inline(paragraph,line,9 if line.startswith('版本：') else None,'66788A' if line.startswith('版本：') else None)
            if line.endswith('：'):paragraph.paragraph_format.keep_with_next=True
        index+=1
    doc.save(OUTPUT)
    print(OUTPUT)
    print('Tables:',len(doc.tables),'Paragraphs:',len(doc.paragraphs))


if __name__=='__main__':build()
