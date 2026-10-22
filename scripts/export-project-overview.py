"""Export the project overview to Word; document tooling is not a runtime dependency.

Run with the Python returned by the Codex workspace dependency loader.
"""
from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'docs/项目总览与实现说明.md'
OUTPUT = SOURCE.with_suffix('.docx')
ASSET = ROOT / 'docs/assets/project-overview-architecture.png'
FONT = '/System/Library/Fonts/Supplemental/Arial Unicode.ttf'


def diagram():
    ASSET.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new('RGB', (2400, 1120), 'white')
    draw = ImageDraw.Draw(image)
    titlefont, smallfont = ImageFont.truetype(FONT, 43), ImageFont.truetype(FONT, 31)
    boxes = [(70, 55, '精确目标群窗口', '已登录  未最小化'),
             (875, 55, '原生辅助功能采集', 'Mac AX 或 Windows UIA'),
             (1680, 55, 'NativeStream', 'JSON 快照  心跳  状态'),
             (1680, 430, '消息标准化与来源校验', '转发详情和 JSONL 在此接入'),
             (875, 430, '本地规则识别', '主题  摘要  截止  活动时间'),
             (70, 430, 'SQLite 持久化', '消息去重  活动关联  版本'),
             (70, 805, 'outbox 通知任务', '新增  更新  取消  待核实'),
             (875, 805, '发送模式与凭据检查', '当前只生成预览'),
             (1680, 805, '邮件预览或 TLS SMTP', '正式发送需明确启用')]
    width, height = 650, 220
    for x, y, title, subtitle in boxes:
        draw.rounded_rectangle((x, y, x + width, y + height), radius=18,
                               fill='#F0F5F9', outline='#7B94A7', width=3)
        draw.text((x + width / 2, y + 75), title, font=titlefont, anchor='mm', fill='#111111')
        draw.text((x + width / 2, y + 150), subtitle, font=smallfont, anchor='mm', fill='#334455')
    arrows = [((720, 165), (875, 165)), ((1525, 165), (1680, 165)),
              ((2005, 275), (2005, 430)), ((1680, 540), (1525, 540)),
              ((875, 540), (720, 540)), ((395, 650), (395, 805)),
              ((720, 915), (875, 915)), ((1525, 915), (1680, 915))]
    for (x1, y1), (x2, y2) in arrows:
        draw.line((x1, y1, x2, y2), fill='#334455', width=7)
        if x2 > x1:
            points = [(x2, y2), (x2 - 22, y2 - 13), (x2 - 22, y2 + 13)]
        elif x2 < x1:
            points = [(x2, y2), (x2 + 22, y2 - 13), (x2 + 22, y2 + 13)]
        else:
            points = [(x2, y2), (x2 - 13, y2 - 22), (x2 + 13, y2 - 22)]
        draw.polygon(points, fill='#334455')
    image.save(ASSET, dpi=(300, 300))


def run_style(run, size=None, bold=False, mono=False):
    run.font.name = 'Menlo' if mono else 'Arial Unicode MS'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def hyperlink(paragraph, label, url, size):
    relation = paragraph.part.relate_to(url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    element = OxmlElement('w:hyperlink'); element.set(qn('r:id'), relation)
    run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for key in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn('w:' + key), 'Arial Unicode MS')
    properties.append(fonts)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '174A72'); properties.append(color)
    fontsize = OxmlElement('w:sz'); fontsize.set(qn('w:val'), str(round(size * 2))); properties.append(fontsize)
    run.append(properties)
    text = OxmlElement('w:t'); text.text = label; run.append(text); element.append(run)
    paragraph._p.append(element)


def inline(paragraph, text, size=10.5):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))', text):
        if not token:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if link:
            hyperlink(paragraph, *link.groups(), size)
            continue
        bold = token.startswith('**') and token.endswith('**')
        mono = token.startswith('`') and token.endswith('`')
        content = token[2:-2] if bold else token[1:-1] if mono else token
        run_style(paragraph.add_run(content), size, bold, mono)


def table(doc, rows):
    count = len(rows[0])
    widths = [4.3, 12.5] if count == 2 else [4.5, 6.2, 6.1]
    if rows[0][0] == '编号':
        widths = [2.0, 14.8]
    result = doc.add_table(rows=0, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + name)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    result._tbl.tblPr.append(borders)
    for number, values in enumerate(rows):
        row = result.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if number == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for name, amount in [('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')]:
                margin = OxmlElement('w:' + name); margin.set(qn('w:w'), amount); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            if number == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E4EDF4'); cell._tc.get_or_add_tcPr().append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.18
            inline(paragraph, value, 9.3)
            if number == 0:
                for run in paragraph.runs:
                    run.bold = True
    after = doc.add_paragraph(); after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.space_before = Pt(0); after.paragraph_format.line_spacing = 0.4


def build():
    diagram()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    for name, size in [('Normal', 10.5), ('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Arial Unicode MS'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        properties = style._element.get_or_add_pPr()
        for border in list(properties.findall(qn('w:pBdr'))):
            properties.remove(border)
        style.paragraph_format.line_spacing = 1.28
        style.paragraph_format.space_after = Pt(6 if name == 'Normal' else 8)
        style.paragraph_format.space_before = Pt(0 if name in ('Normal', 'Title') else 12)
        if name != 'Normal':
            style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_style(footer.add_run('第 '), 8.5)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    run_style(footer.add_run(' 页'), 8.5)
    doc.core_properties.title = '微信讲座跟踪项目总览与实现说明'
    doc.core_properties.subject = '项目来源 模块框架 群消息识别 邮件发送'
    doc.core_properties.author = 'WeChat Lecture Tracker'
    doc.core_properties.created = datetime(2026, 10, 16, 3, 15)
    lines = SOURCE.read_text().splitlines(); index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1; continue
        if line.startswith('```'):
            language = line[3:].strip(); block = []; index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                block.append(lines[index]); index += 1
            if language == 'mermaid':
                paragraph = doc.add_paragraph(); paragraph.paragraph_format.keep_with_next = True
                shape = paragraph.add_run().add_picture(str(ASSET), width=Cm(16.8))
                shape._inline.docPr.set('descr', '核心链路为目标群窗口到原生采集器，再到 JSON 帧、标准化、本地规则、SQLite、通知任务和邮件预览或 SMTP。')
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline(paragraph, '图 1  当前采集 识别 存储与邮件链路', 9)
            else:
                for position, code in enumerate(block):
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.left_indent = Cm(0.2)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.keep_with_next = position < len(block) - 1
                    run_style(paragraph.add_run(code), 9, mono=True)
            index += 1; continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                parts = [part.strip() for part in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', part.replace(' ', '')) for part in parts):
                    rows.append(parts)
                index += 1
            table(doc, rows); continue
        heading = re.match(r'^(#{1,3}) (.*)', line)
        if heading:
            depth, title = len(heading[1]), heading[2]
            paragraph = doc.add_paragraph(title, style='Title' if depth == 1 else 'Heading ' + str(depth - 1))
        elif line.startswith('- '):
            paragraph = doc.add_paragraph(style='List Bullet'); inline(paragraph, line[2:])
        else:
            paragraph = doc.add_paragraph(); inline(paragraph, line)
        index += 1
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
